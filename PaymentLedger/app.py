import os
import re
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
import pandas as pd
import pytesseract

from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    session,
    send_file,
    flash,
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from werkzeug.utils import secure_filename


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

# Render/container-safe directories.
# These are created automatically when the app starts.
UPLOADS = BASE_DIR / "uploads"
DATABASE = BASE_DIR / "database"

UPLOADS.mkdir(parents=True, exist_ok=True)
DATABASE.mkdir(parents=True, exist_ok=True)

EXCEL_FILE = DATABASE / "PaymentLedger.xlsx"


# ============================================================
# APP CONFIG
# ============================================================

app = Flask(__name__)

# For testing only. For real deployment, set FLASK_SECRET_KEY
# as an environment variable in Render.
app.secret_key = os.getenv(
    "FLASK_SECRET_KEY",
    "paymentledger-testing-secret-change-later"
)

OWNER_USERNAME = os.getenv("OWNER_USERNAME", "owner")
OWNER_PASSWORD = os.getenv("OWNER_PASSWORD", "change-this-password")

ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}

COLUMNS = [
    "Record ID",
    "Date",
    "Time",
    "Status",
    "Paid To",
    "UPI ID",
    "Amount",
    "Transaction ID",
    "Debited From",
    "UTR",
    "What Did You Purchase?",
    "Created At",
]


# ============================================================
# DATABASE INITIALIZATION
# ============================================================

def ensure_database():
    """
    Create the Excel database if it does not exist.
    This also protects the app if the database directory
    was not included in the deployment.
    """

    DATABASE.mkdir(parents=True, exist_ok=True)

    if not EXCEL_FILE.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Payments"

        ws.append(COLUMNS)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        ws.freeze_panes = "A2"

        wb.save(EXCEL_FILE)


# ============================================================
# FILE VALIDATION
# ============================================================

def allowed_file(filename):
    return (
        bool(filename)
        and "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


# ============================================================
# OCR
# ============================================================

def extract_ocr(image_path):
    """
    Read payment screenshot and extract OCR text using
    OpenCV + Tesseract.
    """

    image = cv2.imread(str(image_path))

    if image is None:
        raise ValueError("Could not read the uploaded image.")

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)

    # Improve OCR quality
    gray = cv2.resize(
        gray,
        None,
        fx=2,
        fy=2,
        interpolation=cv2.INTER_CUBIC
    )

    gray = cv2.normalize(
        gray,
        None,
        0,
        255,
        cv2.NORM_MINMAX
    )

    gray = cv2.GaussianBlur(
        gray,
        (3, 3),
        0
    )

    data = pytesseract.image_to_data(
        gray,
        config="--oem 3 --psm 6",
        lang="eng",
        output_type=pytesseract.Output.DATAFRAME
    )

    data = data.dropna(subset=["text"])

    data["text"] = (
        data["text"]
        .astype(str)
        .str.strip()
    )

    data = data[
        data["text"] != ""
    ].reset_index(drop=True)

    return data


# ============================================================
# PAYMENT DATA EXTRACTION
# ============================================================

def extract_payment_data(ocr_data):

    words = ocr_data["text"].tolist()

    text = " ".join(words)

    result = {
        "Date": "",
        "Time": "",
        "Status": "",
        "Paid To": "",
        "UPI ID": "",
        "Amount": "",
        "Transaction ID": "",
        "Debited From": "",
        "UTR": "",
    }

    # --------------------------------------------------------
    # DATE
    # --------------------------------------------------------

    date_match = re.search(
        r"\b(\d{1,2})\s+"
        r"(January|February|March|April|May|June|July|August|"
        r"September|October|November|December)\s+"
        r"(\d{4})\b",
        text,
        re.IGNORECASE
    )

    if date_match:
        result["Date"] = " ".join(date_match.groups())


    # --------------------------------------------------------
    # TIME
    # --------------------------------------------------------

    time_match = re.search(
        r"\b(\d{1,2}:\d{2})\b",
        text
    )

    if time_match:
        result["Time"] = time_match.group(1)


    # --------------------------------------------------------
    # STATUS
    # --------------------------------------------------------

    status_match = re.search(
        r"\b(successful|failed|pending)\b",
        text,
        re.IGNORECASE
    )

    if status_match:
        result["Status"] = status_match.group(1).title()


    # --------------------------------------------------------
    # UPI ID
    # --------------------------------------------------------

    upi_match = re.search(
        r"\b[A-Za-z0-9._-]+@[A-Za-z0-9._-]+\b",
        text
    )

    if upi_match:
        result["UPI ID"] = upi_match.group(0)


    # --------------------------------------------------------
    # TRANSACTION ID
    # --------------------------------------------------------

    transaction_match = re.search(
        r"\bT[A-Za-z0-9]{10,}\b",
        text
    )

    if transaction_match:
        result["Transaction ID"] = transaction_match.group(0)


    # --------------------------------------------------------
    # UTR
    # --------------------------------------------------------

    utr_match = re.search(
        r"\bUTR[:\s-]*([0-9]{8,})\b",
        text,
        re.IGNORECASE
    )

    if utr_match:
        result["UTR"] = utr_match.group(1)


    # --------------------------------------------------------
    # AMOUNT
    #
    # Payment screenshots may cause OCR to read ₹ as %
    # Example:
    #
    # %3
    #
    # becomes:
    #
    # Rs 3
    # --------------------------------------------------------

    amount_match = re.search(
        r"(?:₹|Rs\.?|INR|%)"
        r"[\s:]*"
        r"([\d,]+(?:\.\d{1,2})?)",
        text,
        re.IGNORECASE
    )

    if amount_match:
        amount_value = (
            amount_match.group(1)
            .replace(",", "")
        )

        result["Amount"] = f"Rs {amount_value}"


    # --------------------------------------------------------
    # PAID TO
    # --------------------------------------------------------

    paid_to = []

    try:

        paid_index = next(
            i for i, word in enumerate(words)
            if word.lower() == "paid"
        )

        if (
            paid_index + 1 < len(words)
            and words[paid_index + 1].lower() == "to"
        ):

            stop_words = {
                "payment",
                "details",
                "verified",
                "transaction",
                "id",
                "debited",
                "from",
                "utr",
            }

            for word in words[paid_index + 2:]:

                lower_word = word.lower()

                if lower_word in stop_words:
                    break

                if "@" in word:
                    break

                if re.match(
                    r"^(?:₹|Rs|INR|%)",
                    word,
                    re.IGNORECASE
                ):
                    break

                paid_to.append(word)

                if len(paid_to) >= 5:
                    break

    except StopIteration:
        pass

    if paid_to:
        result["Paid To"] = " ".join(paid_to)


    # --------------------------------------------------------
    # DEBITED FROM
    #
    # Preserve the OCR pieces before/around the masked account.
    #
    # Example:
    #
    # «@) XXXXXXXXXX5300
    #
    # --------------------------------------------------------

    try:

        debited_index = next(
            i for i, word in enumerate(words)
            if word.lower() == "debited"
        )

        if (
            debited_index + 1 < len(words)
            and words[debited_index + 1].lower() == "from"
        ):

            pieces = []

            for word in words[debited_index + 2:]:

                # Stop at the next major section.
                if word.lower() in {
                    "payment",
                    "details",
                    "verified",
                    "transaction",
                    "utr",
                }:
                    break

                pieces.append(word)

                # Masked bank/card/account number.
                if (
                    re.search(r"[Xx*]{4,}\d{2,}$", word)
                    or re.search(r"\d{4}$", word)
                ):
                    break

                # Prevent unrelated OCR text from being included.
                if len(pieces) >= 3:
                    break

            result["Debited From"] = " ".join(
                pieces
            ).strip()

    except StopIteration:
        pass


    return result


# ============================================================
# RECORD ID
# ============================================================

def next_record_id(df):

    if df.empty or "Record ID" not in df.columns:
        return 1

    ids = pd.to_numeric(
        df["Record ID"],
        errors="coerce"
    ).dropna()

    if ids.empty:
        return 1

    return int(ids.max()) + 1


# ============================================================
# SAVE PAYMENT TO EXCEL
# ============================================================

def save_payment(payment):

    ensure_database()

    try:
        df = pd.read_excel(
            EXCEL_FILE,
            sheet_name="Payments"
        )
    except Exception:
        df = pd.DataFrame(
            columns=COLUMNS
        )

    # Make sure all expected columns exist.
    for column in COLUMNS:
        if column not in df.columns:
            df[column] = ""

    record_id = next_record_id(df)

    payment["Record ID"] = record_id

    payment["Created At"] = datetime.now().strftime(
        "%Y-%m-%d %H:%M:%S"
    )

    new_row = {
        column: payment.get(column, "")
        for column in COLUMNS
    }

    new_df = pd.DataFrame(
        [new_row]
    )

    df = pd.concat(
        [
            df[COLUMNS],
            new_df
        ],
        ignore_index=True
    )

    # Write Excel
    with pd.ExcelWriter(
        EXCEL_FILE,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            sheet_name="Payments",
            index=False
        )

    # Format Excel
    wb = load_workbook(EXCEL_FILE)

    ws = wb["Payments"]

    ws.freeze_panes = "A2"

    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    widths = {
        "A": 12,
        "B": 18,
        "C": 12,
        "D": 14,
        "E": 25,
        "F": 32,
        "G": 15,
        "H": 30,
        "I": 30,
        "J": 22,
        "K": 30,
        "L": 22,
    }

    for column, width in widths.items():

        ws.column_dimensions[
            column
        ].width = width

    wb.save(EXCEL_FILE)

    return record_id


# ============================================================
# USER VIEW
# ============================================================

@app.route("/", methods=["GET", "POST"])
def user():

    if request.method == "POST":

        f = request.files.get(
            "payment_image"
        )

        purchase = request.form.get(
            "purchase",
            ""
        ).strip()

        # ----------------------------------------------------
        # Validate image
        # ----------------------------------------------------

        if not f or not f.filename:

            flash(
                "Please upload a payment screenshot.",
                "error"
            )

            return redirect(
                url_for("user")
            )

        if not allowed_file(f.filename):

            flash(
                "Please upload PNG, JPG, JPEG or WEBP.",
                "error"
            )

            return redirect(
                url_for("user")
            )

        # ----------------------------------------------------
        # Validate purchase
        # ----------------------------------------------------

        if not purchase:

            flash(
                "Please enter what you purchased.",
                "error"
            )

            return redirect(
                url_for("user")
            )

        # ----------------------------------------------------
        # IMPORTANT FIX
        #
        # Render was failing because /app/uploads did not exist.
        # Create the directory immediately before saving.
        # ----------------------------------------------------

        UPLOADS.mkdir(
            parents=True,
            exist_ok=True
        )

        filename = secure_filename(
            f.filename
        )

        timestamp = datetime.now().strftime(
            "%Y%m%d%H%M%S%f"
        )

        path = UPLOADS / (
            timestamp
            + "_"
            + filename
        )

        try:

            # Save screenshot
            f.save(path)

            # OCR
            ocr_data = extract_ocr(
                path
            )

            # Extract payment details
            payment = extract_payment_data(
                ocr_data
            )

            # Manual purchase field
            payment[
                "What Did You Purchase?"
            ] = purchase

            # Save to Excel
            record_id = save_payment(
                payment
            )

        except Exception as exc:

            app.logger.exception(
                "Payment submission failed"
            )

            flash(
                f"Unable to process payment: {exc}",
                "error"
            )

            return redirect(
                url_for("user")
            )

        finally:

            # Remove temporary uploaded image.
            try:
                if path.exists():
                    path.unlink()
            except Exception:
                pass

        return render_template(
            "user.html",
            submitted=True,
            record_id=record_id
        )

    return render_template(
        "user.html",
        submitted=False
    )


# ============================================================
# OWNER LOGIN
# ============================================================

@app.route(
    "/owner/login",
    methods=["GET", "POST"]
)
def owner_login():

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        )

        password = request.form.get(
            "password",
            ""
        )

        if (
            username == OWNER_USERNAME
            and password == OWNER_PASSWORD
        ):

            session["owner_logged_in"] = True

            return redirect(
                url_for("owner_view")
            )

        flash(
            "Invalid owner credentials.",
            "error"
        )

    return render_template(
        "login.html"
    )


# ============================================================
# OWNER DASHBOARD
# ============================================================

@app.route("/owner")
def owner_view():

    if not session.get(
        "owner_logged_in"
    ):

        return redirect(
            url_for("owner_login")
        )

    ensure_database()

    try:

        df = pd.read_excel(
            EXCEL_FILE,
            sheet_name="Payments"
        )

    except Exception:

        df = pd.DataFrame(
            columns=COLUMNS
        )

    if df.empty:

        records = []
        total = 0

    else:

        df = df.fillna("")

        records = df.to_dict(
            "records"
        )

        amounts = (
            df["Amount"]
            .astype(str)
            .str.replace(
                r"[^0-9.]",
                "",
                regex=True
            )
        )

        amounts = pd.to_numeric(
            amounts,
            errors="coerce"
        ).fillna(0)

        total = float(
            amounts.sum()
        )

    return render_template(
        "owner.html",
        records=records,
        total=f"{total:,.2f}",
        count=len(records)
    )


# ============================================================
# OWNER LOGOUT
# ============================================================

@app.route("/owner/logout")
def owner_logout():

    session.clear()

    return redirect(
        url_for("owner_login")
    )


# ============================================================
# DOWNLOAD EXCEL
# ============================================================

@app.route("/owner/download")
def download_excel():

    if not session.get(
        "owner_logged_in"
    ):

        return redirect(
            url_for("owner_login")
        )

    ensure_database()

    return send_file(
        EXCEL_FILE,
        as_attachment=True,
        download_name="PaymentLedger.xlsx"
    )


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route("/health")
def health():

    return {
        "status": "ok",
        "app": "PaymentLedger"
    }


# ============================================================
# START APP
# ============================================================

if __name__ == "__main__":

    ensure_database()

    port = int(
        os.getenv(
            "PORT",
            "8501"
        )
    )

    app.run(
        host="0.0.0.0",
        port=port,
        debug=False
    )
